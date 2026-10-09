"""Build a formatted Excel workbook from classified Clutch company rows."""
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

OUTREACH_COLS = [
    "Date Added", "Platform", "Agency Name", "Country", "City", "Website",
    "Profile Link", "Main Services", "Secondary Services", "Team Size",
    "Hourly Rate (USD)", "Min Project Size (USD)", "Industries Served",
    "DM Name", "DM Title", "Email", "LinkedIn Profile", "Contact Form Link",
    "Signs They May Outsource", "Qual Score", "Priority", "First Outreach Date",
    "Follow-Up 1 Date", "Follow-Up 2 Date", "Response Status", "Next Action", "Remarks",
]
DAILY_LOG_COLS = [
    "Log Date", "Researcher", "Researched", "Qualified", "DMs Found", "Emails Sent",
    "LinkedIn Sent", "Follow-Ups", "Replies", "Positive", "Calls Sched.", "Notes / Lessons",
]
OOS_COLS = [
    "Date Added", "Platform", "Agency Name", "Country", "City", "Website",
    "Profile Link", "Main Services", "Team Size", "Hourly Rate (USD)",
    "Min Project Size (USD)", "Clutch Rating", "Out of Scope Reason", "Details", "Remarks",
]

_NAVY = PatternFill("solid", fgColor="1F4E79")
_GREEN = PatternFill("solid", fgColor="375623")
_BLUE = PatternFill("solid", fgColor="2E75B6")
_RED = PatternFill("solid", fgColor="843C0C")
_WHITE = Font(color="FFFFFF", bold=True)

_TRACKER_WIDTHS = {
    "Date Added": 14, "Platform": 10, "Agency Name": 30, "Country": 15, "City": 18,
    "Website": 32, "Profile Link": 32, "Main Services": 42, "Secondary Services": 36,
    "Team Size": 12, "Hourly Rate (USD)": 18, "Min Project Size (USD)": 20,
    "Industries Served": 25, "DM Name": 20, "DM Title": 20, "Email": 25,
    "LinkedIn Profile": 30, "Contact Form Link": 30, "Signs They May Outsource": 45,
    "Qual Score": 13, "Priority": 18, "First Outreach Date": 18, "Follow-Up 1 Date": 18,
    "Follow-Up 2 Date": 18, "Response Status": 18, "Next Action": 25, "Remarks": 50,
    "Log Date": 14, "Researcher": 15, "Researched": 13, "Qualified": 13, "DMs Found": 12,
    "Emails Sent": 13, "LinkedIn Sent": 14, "Follow-Ups": 12, "Replies": 10, "Positive": 10,
    "Calls Sched.": 13, "Notes / Lessons": 30, "Also in .NET Tracker?": 22,
}
_OOS_WIDTHS = {
    "Date Added": 14, "Platform": 10, "Agency Name": 30, "Country": 15, "City": 18,
    "Website": 32, "Profile Link": 32, "Main Services": 42, "Team Size": 12,
    "Hourly Rate (USD)": 18, "Min Project Size (USD)": 20, "Clutch Rating": 14,
    "Out of Scope Reason": 45, "Details": 50, "Remarks": 40, "Also in .NET Tracker?": 22,
}


def _apply_header(cell, text, fill, font=None):
    cell.value = text
    cell.fill = fill
    cell.font = font or _WHITE
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _build_tracker_sheet(wb: Workbook, rows: list, today: str, cross_refs: Optional[dict] = None) -> None:
    ws = wb.create_sheet("Agency Tracker")
    extra_cols = list(cross_refs.keys()) if cross_refs else []
    all_cols = OUTREACH_COLS + DAILY_LOG_COLS + extra_cols
    n_outreach, n_log = len(OUTREACH_COLS), len(DAILY_LOG_COLS)

    _apply_header(ws.cell(1, 1), "OUTREACH PIPELINE", _NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_outreach)
    log_start, log_end = n_outreach + 1, n_outreach + n_log
    _apply_header(ws.cell(1, log_start), "DAILY PROGRESS LOG", _GREEN)
    ws.merge_cells(start_row=1, start_column=log_start, end_row=1, end_column=log_end)
    if extra_cols:
        start, end = log_end + 1, log_end + len(extra_cols)
        _apply_header(ws.cell(1, start), "CROSS-REFERENCE", _NAVY)
        if len(extra_cols) > 1:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    ws.row_dimensions[1].height = 22

    for idx, col in enumerate(all_cols, 1):
        _apply_header(ws.cell(2, idx), col, _GREEN if col in DAILY_LOG_COLS else _NAVY)
    ws.row_dimensions[2].height = 36

    for raw, result in rows:
        blank = ""
        name = raw.get("agency_name", "")
        row_data = {
            "Date Added": today, "Platform": raw.get("platform", "Clutch"),
            "Agency Name": name, "Country": raw.get("country", ""), "City": raw.get("city", ""),
            "Website": raw.get("website", ""), "Profile Link": raw.get("profile_link", ""),
            "Main Services": ", ".join(raw.get("main_services") or []),
            "Secondary Services": ", ".join(raw.get("secondary_services") or []),
            "Team Size": raw.get("team_size_raw", ""), "Hourly Rate (USD)": raw.get("hourly_rate_raw", ""),
            "Min Project Size (USD)": raw.get("min_project_size_raw", ""), "Industries Served": blank,
            "DM Name": blank, "DM Title": blank, "Email": blank, "LinkedIn Profile": blank,
            "Contact Form Link": blank, "Signs They May Outsource": result.get("signs_they_may_outsource", ""),
            "Qual Score": f"{result.get('qual_score', '')} / 10", "Priority": result.get("priority", ""),
            "First Outreach Date": blank, "Follow-Up 1 Date": blank, "Follow-Up 2 Date": blank,
            "Response Status": "Not Contacted", "Next Action": blank, "Remarks": result.get("remarks", ""),
        }
        for col in DAILY_LOG_COLS:
            row_data[col] = blank
        for col, name_set in (cross_refs or {}).items():
            row_data[col] = "Yes" if name in name_set else "No"
        ws.append([row_data.get(col, blank) for col in all_cols])

    for idx, col in enumerate(all_cols, 1):
        ws.column_dimensions[get_column_letter(idx)].width = _TRACKER_WIDTHS.get(col, 15)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(all_cols))}{ws.max_row}"


def _build_oos_sheet(wb: Workbook, rows: list, today: str, cross_refs: Optional[dict] = None) -> None:
    ws = wb.create_sheet("Out of Scope")
    extra_cols = list(cross_refs.keys()) if cross_refs else []
    all_cols = OOS_COLS + extra_cols
    _apply_header(ws.cell(1, 1), "OUT OF SCOPE COMPANIES", _RED)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(all_cols))
    ws.row_dimensions[1].height = 22
    for idx, col in enumerate(all_cols, 1):
        _apply_header(ws.cell(2, idx), col, _RED)
    ws.row_dimensions[2].height = 30

    for raw, result in rows:
        name = raw.get("agency_name", "")
        row_data = {
            "Date Added": today, "Platform": raw.get("platform", "Clutch"),
            "Agency Name": name, "Country": raw.get("country", ""), "City": raw.get("city", ""),
            "Website": raw.get("website", ""), "Profile Link": raw.get("profile_link", ""),
            "Main Services": ", ".join(raw.get("main_services") or []),
            "Team Size": raw.get("team_size_raw", ""), "Hourly Rate (USD)": raw.get("hourly_rate_raw", ""),
            "Min Project Size (USD)": raw.get("min_project_size_raw", ""), "Clutch Rating": raw.get("rating", ""),
            "Out of Scope Reason": result.get("out_of_scope_reason", ""),
            "Details": result.get("secondary_reasons", ""), "Remarks": result.get("remarks", ""),
        }
        for col, name_set in (cross_refs or {}).items():
            row_data[col] = "Yes" if name in name_set else "No"
        ws.append([row_data.get(col, "") for col in all_cols])

    for idx, col in enumerate(all_cols, 1):
        ws.column_dimensions[get_column_letter(idx)].width = _OOS_WIDTHS.get(col, 15)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(all_cols))}{ws.max_row}"


def _copy_scoring_guide(wb: Workbook, template_path: str) -> None:
    path = Path(template_path)
    if not path.exists():
        return
    src_wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if "Scoring Guide" not in src_wb.sheetnames:
            return
        src_ws = src_wb["Scoring Guide"]
        dst_ws = wb.create_sheet("Scoring Guide")
        from copy import copy
        for row in src_ws.iter_rows():
            for cell in row:
                dst_cell = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    dst_cell.font = copy(cell.font)
                    dst_cell.fill = copy(cell.fill)
                    dst_cell.alignment = copy(cell.alignment)
        for key, dim in src_ws.column_dimensions.items():
            dst_ws.column_dimensions[key].width = dim.width
    finally:
        src_wb.close()


def write_workbook(
    classified_rows: List[tuple],
    template_path: str,
    output_path: str,
    cross_refs: Optional[dict] = None,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    today = pd.Timestamp.today().strftime("%d-%b-%Y")
    tracker_rows = [(raw, res) for raw, res in classified_rows if res.get("scope") == "tracker"]
    oos_rows = [(raw, res) for raw, res in classified_rows if res.get("scope") != "tracker"]
    _build_tracker_sheet(wb, tracker_rows, today, cross_refs)
    _build_oos_sheet(wb, oos_rows, today, cross_refs)
    _copy_scoring_guide(wb, template_path)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
